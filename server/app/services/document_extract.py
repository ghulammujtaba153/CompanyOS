"""
Extract searchable text and tables from uploaded vault files for RAG.

Tables are serialized as Markdown so embeddings keep row/column structure
and Company Brain can answer questions about tabular data.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)

# Keep extracted body bounded for DB + indexing.
MAX_EXTRACT_CHARS = 120_000
MAX_TABLE_ROWS = 200
MAX_TABLE_COLS = 40


@dataclass
class ExtractedTable:
    title: str
    headers: list[str]
    rows: list[list[str]]
    source: str = ""  # e.g. sheet name, "page 2"

    def to_markdown(self) -> str:
        headers = self.headers or [f"Col{i + 1}" for i in range(max((len(r) for r in self.rows), default=0))]
        if not headers and not self.rows:
            return ""
        width = min(MAX_TABLE_COLS, max(len(headers), max((len(r) for r in self.rows), default=0)))
        headers = [(headers[i] if i < len(headers) else f"Col{i + 1}").strip() or f"Col{i + 1}" for i in range(width)]
        lines = [
            f"### Table: {self.title}" + (f" ({self.source})" if self.source else ""),
            "| " + " | ".join(_cell(h) for h in headers) + " |",
            "| " + " | ".join("---" for _ in headers) + " |",
        ]
        for row in self.rows[:MAX_TABLE_ROWS]:
            cells = [_cell(row[i] if i < len(row) else "") for i in range(width)]
            lines.append("| " + " | ".join(cells) + " |")
        if len(self.rows) > MAX_TABLE_ROWS:
            lines.append(f"_… {len(self.rows) - MAX_TABLE_ROWS} more rows omitted_")
        return "\n".join(lines)


@dataclass
class ExtractResult:
    text: str = ""
    tables: list[ExtractedTable] = field(default_factory=list)
    format: str = ""
    warnings: list[str] = field(default_factory=list)

    @property
    def table_markdown(self) -> str:
        parts = [t.to_markdown() for t in self.tables if t.to_markdown()]
        return "\n\n".join(parts)

    def combined_body(self, *, user_notes: str = "", filename: str = "") -> str:
        sections: list[str] = []
        notes = (user_notes or "").strip()
        if notes:
            sections.append(notes)
        if self.text.strip():
            sections.append(self.text.strip())
        tables = self.table_markdown
        if tables:
            sections.append("## Tables\n\n" + tables)
        if not sections:
            return f"Uploaded document: {filename}".strip() if filename else ""
        body = "\n\n".join(sections)
        if len(body) > MAX_EXTRACT_CHARS:
            body = body[: MAX_EXTRACT_CHARS - 20] + "\n\n…[truncated]"
        return body


def _cell(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\s+", " ", text).strip()
    return text.replace("|", "\\|")[:200]


def _ext(filename: str | None) -> str:
    return Path(filename or "").suffix.lower().lstrip(".")


def extract_bytes(raw: bytes, *, filename: str | None = None, content_type: str | None = None) -> ExtractResult:
    """Best-effort extract text + tables. Never raises for unsupported types."""
    ext = _ext(filename)
    ctype = (content_type or "").split(";")[0].strip().lower()
    if not ext:
        if "csv" in ctype:
            ext = "csv"
        elif "spreadsheet" in ctype or "excel" in ctype:
            ext = "xlsx"
        elif "pdf" in ctype:
            ext = "pdf"
        elif "word" in ctype or "officedocument.wordprocessingml" in ctype:
            ext = "docx"
        elif ctype.startswith("text/"):
            ext = "txt"

    try:
        if ext in {"txt", "md", "markdown", "log"}:
            return _extract_plain(raw, ext)
        if ext in {"csv", "tsv"}:
            return _extract_csv(raw, ext)
        if ext in {"xlsx", "xlsm"}:
            return _extract_xlsx(raw, filename or "workbook.xlsx")
        if ext == "pdf":
            return _extract_pdf(raw)
        if ext in {"docx", "doc"}:
            return _extract_docx(raw)
    except Exception as exc:
        logger.warning("Document extract failed for %s: %s", filename, exc)
        return ExtractResult(format=ext or "unknown", warnings=[f"Extract failed: {exc}"])

    return ExtractResult(
        format=ext or "unknown",
        warnings=[f"No text extractor for .{ext or 'bin'} — add notes for Company Brain."],
    )


def _extract_plain(raw: bytes, ext: str) -> ExtractResult:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            text = ""
    else:
        text = raw.decode("utf-8", errors="replace")
    tables = _tables_from_markdown(text) if ext in {"md", "markdown"} else []
    return ExtractResult(text=text.strip(), tables=tables, format=ext)


def _extract_csv(raw: bytes, ext: str) -> ExtractResult:
    text = raw.decode("utf-8-sig", errors="replace")
    delimiter = "\t" if ext == "tsv" else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    matrix = [list(row) for row in reader if any(str(c).strip() for c in row)]
    if not matrix:
        return ExtractResult(text=text.strip(), format=ext)
    headers = [str(c) for c in matrix[0]]
    rows = [[str(c) for c in row] for row in matrix[1:]]
    table = ExtractedTable(title="Spreadsheet", headers=headers, rows=rows, source=ext.upper())
    return ExtractResult(
        text=f"CSV/TSV export with {len(rows)} data row(s) and {len(headers)} column(s).",
        tables=[table],
        format=ext,
    )


def _extract_xlsx(raw: bytes, filename: str) -> ExtractResult:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ExtractResult(
            format="xlsx",
            warnings=["openpyxl not installed — cannot read Excel tables."],
        )
    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    tables: list[ExtractedTable] = []
    text_bits: list[str] = [f"Workbook: {filename}"]
    for sheet in wb.worksheets:
        matrix: list[list[str]] = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= MAX_TABLE_ROWS + 1:
                break
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                matrix.append(cells[:MAX_TABLE_COLS])
        if not matrix:
            continue
        headers = matrix[0]
        rows = matrix[1:]
        tables.append(
            ExtractedTable(
                title=sheet.title,
                headers=headers,
                rows=rows,
                source=f"sheet:{sheet.title}",
            )
        )
        text_bits.append(f"Sheet “{sheet.title}”: {len(rows)} rows × {len(headers)} cols.")
    wb.close()
    return ExtractResult(text="\n".join(text_bits), tables=tables, format="xlsx")


def _extract_pdf(raw: bytes) -> ExtractResult:
    text_parts: list[str] = []
    tables: list[ExtractedTable] = []
    warnings: list[str] = []

    # Prefer pdfplumber for table detection; fall back to pypdf text-only.
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            for page_idx, page in enumerate(pdf.pages, start=1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    text_parts.append(page_text)
                try:
                    raw_tables = page.extract_tables() or []
                except Exception:
                    raw_tables = []
                for t_idx, matrix in enumerate(raw_tables, start=1):
                    cleaned = [
                        ["" if c is None else str(c) for c in row]
                        for row in (matrix or [])
                        if row and any(str(c or "").strip() for c in row)
                    ]
                    if len(cleaned) < 2:
                        continue
                    tables.append(
                        ExtractedTable(
                            title=f"PDF table {t_idx}",
                            headers=cleaned[0],
                            rows=cleaned[1:],
                            source=f"page {page_idx}",
                        )
                    )
        return ExtractResult(text="\n\n".join(text_parts).strip(), tables=tables, format="pdf")
    except ImportError:
        warnings.append("pdfplumber not installed — PDF tables not detected.")
    except Exception as exc:
        warnings.append(f"pdfplumber failed: {exc}")

    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        for page in reader.pages:
            page_text = (page.extract_text() or "").strip()
            if page_text:
                text_parts.append(page_text)
        return ExtractResult(
            text="\n\n".join(text_parts).strip(),
            tables=tables,
            format="pdf",
            warnings=warnings,
        )
    except ImportError:
        return ExtractResult(format="pdf", warnings=warnings + ["pypdf not installed."])
    except Exception as exc:
        return ExtractResult(format="pdf", warnings=warnings + [str(exc)])


def _extract_docx(raw: bytes) -> ExtractResult:
    try:
        from docx import Document
    except ImportError:
        return ExtractResult(format="docx", warnings=["python-docx not installed."])

    doc = Document(io.BytesIO(raw))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    tables: list[ExtractedTable] = []
    for idx, table in enumerate(doc.tables, start=1):
        matrix: list[list[str]] = []
        for row in table.rows[: MAX_TABLE_ROWS + 1]:
            cells = [cell.text.strip() for cell in row.cells[:MAX_TABLE_COLS]]
            if any(cells):
                matrix.append(cells)
        if len(matrix) < 1:
            continue
        if len(matrix) == 1:
            tables.append(
                ExtractedTable(title=f"Word table {idx}", headers=matrix[0], rows=[], source="docx")
            )
        else:
            tables.append(
                ExtractedTable(
                    title=f"Word table {idx}",
                    headers=matrix[0],
                    rows=matrix[1:],
                    source="docx",
                )
            )
    return ExtractResult(text="\n\n".join(paragraphs), tables=tables, format="docx")


def _tables_from_markdown(text: str) -> list[ExtractedTable]:
    """Parse simple GitHub-flavored markdown tables already in the file."""
    tables: list[ExtractedTable] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines) - 1:
        if "|" in lines[i] and re.match(r"^\s*\|?[\s:-|]+\|?\s*$", lines[i + 1] or ""):
            header = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and "|" in lines[i]:
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            tables.append(ExtractedTable(title=f"Markdown table {len(tables) + 1}", headers=header, rows=rows))
            continue
        i += 1
    return tables


def chunk_for_rag(body: str, *, size: int = 900, overlap: int = 120) -> list[dict[str, str]]:
    """
    Split body into chunks. Markdown tables stay intact (or row-batched)
    so embeddings don't smash columns into a single whitespace blob.
    Returns [{text, kind}] where kind is 'text' | 'table'.
    """
    body = (body or "").strip()
    if not body:
        return []

    sections = _split_preserving_tables(body)
    out: list[dict[str, str]] = []
    for section in sections:
        kind = "table" if section.lstrip().startswith("|") or "\n| " in section[:200] else "text"
        if kind == "table":
            out.extend(_chunk_table_block(section, size=size))
        else:
            out.extend(_chunk_prose(section, size=size, overlap=overlap))
    return out or [{"text": body[:size], "kind": "text"}]


def _split_preserving_tables(body: str) -> list[str]:
    lines = body.splitlines()
    sections: list[str] = []
    buf: list[str] = []
    in_table = False

    def flush() -> None:
        nonlocal buf
        text = "\n".join(buf).strip()
        if text:
            sections.append(text)
        buf = []

    i = 0
    while i < len(lines):
        line = lines[i]
        is_table_line = line.strip().startswith("|") and "|" in line.strip()[1:]
        if is_table_line and not in_table:
            flush()
            in_table = True
            # Include a preceding "### Table:" heading if present.
            if sections and sections[-1].lstrip().startswith("### Table"):
                buf.append(sections.pop())
            buf.append(line)
        elif is_table_line and in_table:
            buf.append(line)
        elif in_table and not is_table_line:
            flush()
            in_table = False
            buf.append(line)
        else:
            buf.append(line)
        i += 1
    flush()
    return sections


def _chunk_table_block(block: str, *, size: int) -> list[dict[str, str]]:
    lines = block.splitlines()
    title_lines = [ln for ln in lines if ln.startswith("###")]
    table_lines = [ln for ln in lines if ln.strip().startswith("|")]
    if len(table_lines) < 2:
        return [{"text": block[:size], "kind": "table"}]

    header = table_lines[0]
    sep = table_lines[1] if len(table_lines) > 1 else "| --- |"
    data = table_lines[2:]
    prefix = "\n".join(title_lines + [header, sep])
    chunks: list[dict[str, str]] = []
    batch: list[str] = []
    for row in data:
        candidate = prefix + "\n" + "\n".join(batch + [row])
        if batch and len(candidate) > size:
            chunks.append({"text": prefix + "\n" + "\n".join(batch), "kind": "table"})
            batch = [row]
        else:
            batch.append(row)
    if batch:
        chunks.append({"text": prefix + "\n" + "\n".join(batch), "kind": "table"})
    return chunks or [{"text": block[:size], "kind": "table"}]


def _chunk_prose(text: str, *, size: int, overlap: int) -> list[dict[str, str]]:
    cleaned = re.sub(r"[ \t]+", " ", text)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    if not cleaned:
        return []
    if len(cleaned) <= size:
        return [{"text": cleaned, "kind": "text"}]
    chunks: list[dict[str, str]] = []
    start = 0
    while start < len(cleaned):
        end = min(len(cleaned), start + size)
        # Prefer paragraph / sentence boundaries.
        if end < len(cleaned):
            window = cleaned[start:end]
            cut = max(window.rfind("\n\n"), window.rfind(". "), window.rfind(" "))
            if cut > size * 0.4:
                end = start + cut + 1
        piece = cleaned[start:end].strip()
        if piece:
            chunks.append({"text": piece, "kind": "text"})
        if end >= len(cleaned):
            break
        start = max(0, end - overlap)
    return chunks
